# -*- coding: utf-8 -*-
"""派遣会社リストをExcel出力する"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

JSON_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/haken_companies.json"
OUT_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/労働者派遣_拠点数ランキング_全国TOP500.xlsx"
TOP_N = int(sys.argv[1]) if len(sys.argv) > 1 else 500

HEADERS = ["順位", "事業主名称", "拠点数", "マージン率(平均)", "派遣料金(平均/日)",
           "賃金(平均/日)", "公式HP", "本社所在地", "電話", "記載のある拠点数",
           "許可番号", "許可年月日"]


def main():
    rows = json.load(open(JSON_PATH, encoding="utf-8"))[:TOP_N]

    wb = Workbook()
    ws = wb.active
    ws.title = "派遣会社ランキング"

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="366092")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=r["事業主名称"])
        ws.cell(row=i, column=3, value=r["拠点数"]).number_format = "#,##0"
        c4 = ws.cell(row=i, column=4, value=r.get("マージン率_平均"))
        c4.number_format = '0.0"%"'
        ws.cell(row=i, column=5, value=r.get("派遣料金_平均")).number_format = "#,##0"
        ws.cell(row=i, column=6, value=r.get("賃金_平均")).number_format = "#,##0"
        hp = r.get("公式HP", "")
        if hp:
            c7 = ws.cell(row=i, column=7, value=hp)
            c7.hyperlink = hp
            c7.font = Font(color="0563C1", underline="single", name="Arial")
        else:
            ws.cell(row=i, column=7, value="")
        ws.cell(row=i, column=8, value=r.get("本社所在地", ""))
        ws.cell(row=i, column=9, value=r.get("電話", ""))
        ws.cell(row=i, column=10, value=r.get("記載のある拠点数"))
        ws.cell(row=i, column=11, value=r.get("許可番号", ""))
        ws.cell(row=i, column=12, value=r.get("許可年月日", ""))

    for col, w in zip("ABCDEFGHIJKL", [6, 36, 9, 13, 14, 13, 44, 44, 15, 12, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{len(rows)+1}"
    ws.row_dimensions[1].height = 34

    wb.save(OUT_PATH)
    got = len([r for r in rows if r.get("公式HP")])
    print(f"出力: {OUT_PATH}")
    print(f"  {len(rows)}社 / HP有 {got}社 ({100*got//max(len(rows),1)}%)")


if __name__ == "__main__":
    main()
