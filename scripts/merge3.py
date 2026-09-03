# -*- coding: utf-8 -*-
"""3シート（KBリスト・派遣数Top500・拠点数Top500）を投げ込みリスト用に統合する。

最重要要件: 同一企業（同一の問い合わせ窓口）が2行にならないこと。
  キー1: 許可番号（同一なら同一企業）
  キー2: HPドメイン（別法人でも同一サイト＝同一窓口。グループとして1行に統合）
  キー3: 正規化名+法人格（許可番号なし行の吸収。法人格が違えば別会社として分離）

数値はシートごとに意味が確定しているので、同一許可番号の統合時は
戦略シート側（数値が豊富）を正とし、KB側からはサイト名・備考を持ち込む。
ドメイン統合時は代表1社の数値のみ表示し、他社の数値は合算しない（別法人のため）。
"""
import json
import os
import re
import sys
from collections import OrderedDict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from jinzai_common import normalize_name

BASE = "/Users/masakatsu/Desktop/kyujin_box_rpa"
OUT = os.path.join(BASE, "投げ込みリスト_統合版.xlsx")

sys.path.insert(0, os.path.join(BASE, "scripts"))
from kb_pipeline import company_variants, clean_service, CORP_PAT


def corp_form(name):
    """法人格の種別（同名判定で法人格違いを別会社として扱うため）"""
    s = str(name)
    for f in ("株式会社", "有限会社", "合同会社", "合資会社", "協同組合"):
        if f in s:
            return f
    return ""


def domain(u):
    if not u:
        return ""
    m = re.match(r"https?://([^/]+)", str(u).lower())
    d = m.group(1) if m else ""
    return d[4:] if d.startswith("www.") else d


def to_num(x):
    if x is None or x == "":
        return None
    try:
        return float(x) if isinstance(x, float) or "." in str(x) else int(x)
    except (ValueError, TypeError):
        return None


def main():
    reso = json.load(open(os.path.join(BASE, ".kb_resolution.json"), encoding="utf-8"))
    db = {r["許可番号"]: r for r in
          json.load(open(os.path.join(BASE, "haken_full.json"), encoding="utf-8"))}

    wb_kb = openpyxl.load_workbook("/Users/masakatsu/Downloads/派遣KBリスト.xlsx", read_only=True)
    kb_raws = [str(r[0]).strip() for r in wb_kb.worksheets[0].iter_rows(min_row=2, values_only=True)
               if r[0] and str(r[0]).strip()]
    kb_uniq = list(OrderedDict.fromkeys(kb_raws))

    wb_st = openpyxl.load_workbook(os.path.join(BASE, "戦略_HP補完済み.xlsx"), read_only=True)
    strat = {}
    for ws in wb_st.worksheets:
        header = [str(c) if c else "" for c in
                  next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        strat[ws.title] = [dict(zip(header, r)) for r in
                           ws.iter_rows(min_row=2, values_only=True) if any(r)]

    # ---------------- 1) 許可番号単位で企業ユニットを構築 ----------------
    units = {}  # lic -> unit

    def unit_for(lic):
        if lic not in units:
            rec = db.get(lic, {})
            units[lic] = {
                "出典": [], "サイト名": [], "会社名": rec.get("事業主名称", ""),
                "会社HP": "", "派遣番号": lic,
                "派遣労働者数": None, "公開拠点数": None, "拠点数": None,
                "派遣先件数": None, "マージン率": None, "派遣料金": None, "賃金": None,
                "得意職種": rec.get("得意職種", ""),
                "本社所在地": rec.get("本社所在地", ""), "電話": rec.get("電話", ""),
                "許可年月日": rec.get("許可年月日", ""),
                "派遣数順位": None, "拠点数順位": None, "備考": [],
            }
        return units[lic]

    for r in strat["派遣数Top500"]:
        u = unit_for(r["許可番号"])
        u["出典"].append("派遣数Top500")
        u["会社名"] = r["事業主名称"]
        u["会社HP"] = r.get("公式HP") or u["会社HP"]
        u["派遣数順位"] = r.get("順位")
        u["派遣労働者数"] = to_num(r.get("派遣労働者数"))
        u["公開拠点数"] = to_num(r.get("公開拠点数"))
        u["拠点数"] = to_num(r.get("拠点数"))
        u["派遣先件数"] = to_num(r.get("派遣先件数"))
        u["マージン率"] = to_num(r.get("マージン率(平均)"))
        u["派遣料金"] = to_num(r.get("派遣料金(平均/日)"))
        u["賃金"] = to_num(r.get("賃金(平均/日)"))
        u["得意職種"] = r.get("得意職種") or u["得意職種"]
        u["本社所在地"] = r.get("本社所在地") or u["本社所在地"]
        u["電話"] = r.get("電話") or u["電話"]
        u["許可年月日"] = r.get("許可年月日") or u["許可年月日"]

    for r in strat["派遣拠点数Top500"]:
        u = unit_for(r["許可番号"])
        u["出典"].append("拠点数Top500")
        u["会社名"] = u["会社名"] or r["事業主名称"]
        u["会社HP"] = u["会社HP"] or (r.get("公式HP") or "")
        u["拠点数順位"] = r.get("順位")
        if u["拠点数"] is None:
            u["拠点数"] = to_num(r.get("拠点数"))
        if u["公開拠点数"] is None:
            u["公開拠点数"] = to_num(r.get("記載のある拠点数"))
        if u["マージン率"] is None:
            u["マージン率"] = to_num(r.get("マージン率(平均)"))
        if u["派遣料金"] is None:
            u["派遣料金"] = to_num(r.get("派遣料金(平均/日)"))
        if u["賃金"] is None:
            u["賃金"] = to_num(r.get("賃金(平均/日)"))
        u["本社所在地"] = u["本社所在地"] or (r.get("本社所在地") or "")
        u["電話"] = u["電話"] or (r.get("電話") or "")
        u["許可年月日"] = u["許可年月日"] or (r.get("許可年月日") or "")

    for raw in kb_uniq:
        v = reso.get(raw, {})
        if v.get("派遣番号"):
            u = unit_for(v["派遣番号"])
            if "KBリスト" not in u["出典"]:
                u["出典"].append("KBリスト")
            u["サイト名"].append(raw)
            u["会社名"] = u["会社名"] or v.get("会社名", "")
            u["会社HP"] = u["会社HP"] or v.get("会社HP", "")

    # DBから数値の補完（戦略シートに無い企業＝KB由来）
    for lic, u in units.items():
        rec = db.get(lic, {})
        if u["派遣労働者数"] is None:
            u["派遣労働者数"] = rec.get("派遣労働者数")
        if u["拠点数"] is None:
            u["拠点数"] = rec.get("拠点数")
        if u["公開拠点数"] is None:
            u["公開拠点数"] = rec.get("公開拠点数")
        if u["派遣先件数"] is None:
            u["派遣先件数"] = rec.get("派遣先件数")
        if u["マージン率"] is None:
            u["マージン率"] = rec.get("マージン率_平均")
        if u["派遣料金"] is None:
            u["派遣料金"] = rec.get("派遣料金_平均")
        if u["賃金"] is None:
            u["賃金"] = rec.get("賃金_平均")

    print(f"許可番号ユニット: {len(units)}社")

    # ---------------- 2) 同一ドメイン統合（同一問い合わせ窓口） ----------------
    by_dom = {}
    for lic, u in units.items():
        d = domain(u["会社HP"])
        if d:
            by_dom.setdefault(d, []).append(lic)

    merged_away = set()
    for d, lics in by_dom.items():
        if len(lics) < 2:
            continue
        # 代表 = 派遣労働者数が最大（None は最小扱い）
        lics_sorted = sorted(lics, key=lambda l: -(units[l]["派遣労働者数"] or -1))
        primary = units[lics_sorted[0]]
        others = [units[l] for l in lics_sorted[1:]]
        names = "、".join(f"{o['会社名']}({o['派遣番号']})" for o in others)
        primary["備考"].append(f"同一サイトのグループ会社を統合: {names}")
        primary["派遣番号"] = "、".join([primary["派遣番号"]] + [o["派遣番号"] for o in others])
        for o in others:
            for s in o["出典"]:
                if s not in primary["出典"]:
                    primary["出典"].append(s)
            primary["サイト名"].extend(o["サイト名"])
            merged_away.update([o["派遣番号"].split("、")[0]])
        print(f"  ドメイン統合 {d}: {len(lics)}社 → 1行 (代表: {primary['会社名'][:20]})")
    for lic in merged_away:
        units.pop(lic, None)
    print(f"ドメイン統合後: {len(units)}社")

    # 2b) ドメイングループと同一ブランドだがHPが無くて統合から漏れた行を吸収
    #     （例: ホットスタッフ行橋/小倉。検証時にHPを空欄化したためドメインでは繋がらないが、
    #       同一フランチャイズ＝同一問い合わせ窓口なので分けると二重投げ込みになる）
    GEO = (r"(東日本|西日本|東海|関西|関東|中部|北陸|九州|四国|中国|北海道|東北|"
           r"新潟|岐阜|品川|行橋|小倉|安城|高松|札幌|仙台|名古屋|大阪|福岡|東京|横浜|神戸|京都)")

    def brand_core(name):
        t = re.sub(r"株式会社|有限会社|合同会社|[\s　]", "", str(name))
        t = re.sub(GEO, "", t)
        return normalize_name(t)

    # 3社以上の大型フランチャイズグループのみ前方一致吸収の対象にする
    # （2社グループまで広げると「日本テクニカル」→「日本テクニカル・サービス」等の
    #   別会社を誤って吸収する恐れがあるため）
    group_units = {}
    for lic, u in units.items():
        for b in u["備考"]:
            if "同一サイトのグループ会社を統合" in b and b.count("派") >= 2:
                group_units[brand_core(u["会社名"])] = lic

    def strip_form(name):
        return normalize_name(re.sub(r"株式会社|有限会社|合同会社|[\s　]", "", str(name)))

    brand_merged = []
    for lic in list(units.keys()):
        u = units[lic]
        if u["会社HP"] or lic in group_units.values():
            continue
        core = None
        stripped = strip_form(u["会社名"])
        for brand in group_units:
            if brand and len(brand) >= 5 and stripped.startswith(brand):
                core = brand
                break
        if core:
            g = units[group_units[core]]
            g["備考"].append(f"同一ブランドのため統合: {u['会社名']}({u['派遣番号']})")
            g["派遣番号"] += "、" + u["派遣番号"]
            for s_ in u["出典"]:
                if s_ not in g["出典"]:
                    g["出典"].append(s_)
            g["サイト名"].extend(u["サイト名"])
            brand_merged.append(lic)
            print(f"  ブランド統合: {u['会社名']} → {g['会社名']}")
    for lic in brand_merged:
        units.pop(lic, None)
    print(f"ブランド統合後: {len(units)}社")

    # ---------------- 3) 許可番号なし行の処理 ----------------
    # 会社名(正規化+法人格)→ユニット の索引
    name_idx = {}
    for lic, u in units.items():
        key = (normalize_name(u["会社名"]), corp_form(u["会社名"]))
        name_idx.setdefault(key, []).append(lic)

    extra_rows = []
    absorbed = dup_flagged = 0
    seen_noname = {}
    for raw in kb_uniq:
        v = reso.get(raw, {})
        if v.get("派遣番号"):
            continue
        if CORP_PAT.search(raw):
            cleaned = company_variants(raw)[-1]
            key = (normalize_name(cleaned), corp_form(cleaned))
            hits = name_idx.get(key, [])
            if len(hits) == 1:
                # 表示名が同一企業がリスト内に1社だけ → 二重投げ込み防止のため統合
                u = units[hits[0]]
                if "KBリスト" not in u["出典"]:
                    u["出典"].append("KBリスト")
                u["サイト名"].append(raw)
                if v.get("status") == "ambiguous":
                    u["備考"].append("KB側は同名複数社のため同一社とみなして統合(要確認)")
                absorbed += 1
                continue
            note = v.get("note", "")
            if len(hits) > 1:
                note = (note + " / " if note else "") + "リスト内に同名の許可企業が複数あり要確認"
                dup_flagged += 1
            k2 = key
            if k2 in seen_noname:
                seen_noname[k2]["サイト名"].append(raw)
                continue
            row = {"出典": ["KBリスト"], "サイト名": [raw], "会社名": cleaned,
                   "会社HP": "", "派遣番号": "", "派遣労働者数": None, "公開拠点数": None,
                   "拠点数": None, "派遣先件数": None, "マージン率": None, "派遣料金": None,
                   "賃金": None, "得意職種": "", "本社所在地": "", "電話": "",
                   "許可年月日": "", "派遣数順位": None, "拠点数順位": None,
                   "備考": [note] if note else []}
            seen_noname[k2] = row
            extra_rows.append(row)
        else:
            svc = clean_service(raw)
            k2 = ("SVC:" + normalize_name(svc), "")
            if k2 in seen_noname:
                seen_noname[k2]["サイト名"].append(raw)
                continue
            note = v.get("note", "")
            row = {"出典": ["KBリスト"], "サイト名": [raw], "会社名": "",
                   "会社HP": "", "派遣番号": "", "派遣労働者数": None, "公開拠点数": None,
                   "拠点数": None, "派遣先件数": None, "マージン率": None, "派遣料金": None,
                   "賃金": None, "得意職種": "", "本社所在地": "", "電話": "",
                   "許可年月日": "", "派遣数順位": None, "拠点数順位": None,
                   "備考": [note] if note else []}
            seen_noname[k2] = row
            extra_rows.append(row)

    print(f"許可なし行: 統合吸収{absorbed} / 同名複数フラグ{dup_flagged} / 独立行{len(extra_rows)}")

    # ---------------- 4) 監査: 残存する同一表示名・同一電話 ----------------
    all_rows = list(units.values()) + extra_rows
    from collections import Counter
    name_c = Counter((normalize_name(r["会社名"]), corp_form(r["会社名"]))
                     for r in all_rows if r["会社名"])
    same_names = {k: n for k, n in name_c.items() if n > 1}
    for r in all_rows:
        k = (normalize_name(r["会社名"]), corp_form(r["会社名"]))
        if r["会社名"] and k in same_names:
            r["備考"].append("同名の別会社がリスト内にあり(許可番号で区別)")
    tel_c = Counter(re.sub(r"\D", "", str(r["電話"])) for r in all_rows if r.get("電話"))
    same_tel = {t for t, n in tel_c.items() if n > 1 and len(t) >= 9}
    tel_flag = 0
    for r in all_rows:
        t = re.sub(r"\D", "", str(r.get("電話", "")))
        if t in same_tel:
            r["備考"].append("同一電話番号の企業がリスト内にあり要確認")
            tel_flag += 1
    print(f"監査: 同名別会社 {sum(same_names.values())}行 / 同一電話フラグ {tel_flag}行")

    # ---------------- 5) 出力 ----------------
    all_rows.sort(key=lambda r: (-(r["派遣労働者数"] or -1), -(r["拠点数"] or -1),
                                 r["会社名"] or "～", "、".join(r["サイト名"])))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "投げ込みリスト統合版"
    headers = ["No.", "出典", "サイト名(KB)", "会社名", "会社HP", "派遣番号",
               "派遣労働者数", "公開拠点数", "拠点数", "派遣先件数",
               "マージン率(平均)", "派遣料金(平均/日)", "賃金(平均/日)",
               "得意職種", "本社所在地", "電話", "許可年月日",
               "派遣数順位", "拠点数順位", "備考"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="366092")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(all_rows, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value="、".join(r["出典"]))
        ws.cell(row=i, column=3, value="、".join(OrderedDict.fromkeys(r["サイト名"]))[:300])
        ws.cell(row=i, column=4, value=r["会社名"])
        if r["会社HP"]:
            c5 = ws.cell(row=i, column=5, value=r["会社HP"])
            c5.hyperlink = r["会社HP"]
            c5.font = Font(color="0563C1", underline="single", name="Arial")
        ws.cell(row=i, column=6, value=r["派遣番号"])
        ws.cell(row=i, column=7, value=r["派遣労働者数"])
        ws.cell(row=i, column=8, value=r["公開拠点数"])
        ws.cell(row=i, column=9, value=r["拠点数"])
        ws.cell(row=i, column=10, value=r["派遣先件数"])
        ws.cell(row=i, column=11, value=r["マージン率"])
        ws.cell(row=i, column=12, value=r["派遣料金"])
        ws.cell(row=i, column=13, value=r["賃金"])
        ws.cell(row=i, column=14, value=(r["得意職種"] or "")[:100])
        ws.cell(row=i, column=15, value=r["本社所在地"])
        ws.cell(row=i, column=16, value=r["電話"])
        ws.cell(row=i, column=17, value=r["許可年月日"])
        ws.cell(row=i, column=18, value=r["派遣数順位"])
        ws.cell(row=i, column=19, value=r["拠点数順位"])
        ws.cell(row=i, column=20, value=" / ".join(OrderedDict.fromkeys(r["備考"]))[:250])

    widths = [6, 22, 34, 30, 40, 16, 11, 9, 8, 10, 11, 12, 11, 30, 40, 14, 12, 9, 9, 40]
    for col, w in zip("ABCDEFGHIJKLMNOPQRST", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:T{len(all_rows)+1}"
    ws.row_dimensions[1].height = 30
    wb.save(OUT)

    print(f"\n===== 完了 =====")
    print(f"総行数: {len(all_rows):,}行")
    print(f"  許可あり: {len(units):,}行 / 許可なし: {len(extra_rows):,}行")
    print(f"出力: {OUT}")


if __name__ == "__main__":
    main()
