# -*- coding: utf-8 -*-
"""派遣会社リスト（派遣労働者数×拠点数）をExcel出力する。
シート1: 実数TOP500（HP付き）
シート2: 実数が取れた全社（ソート用の生データ）
"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

JSON_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/haken_full.json"
OUT_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/労働者派遣_派遣数ランキング_全国TOP500.xlsx"

HEADERS = ["順位", "事業主名称", "派遣労働者数", "公開拠点数", "拠点数", "派遣先件数",
           "マージン率(平均)", "派遣料金(平均/日)", "賃金(平均/日)", "公式HP",
           "本社所在地", "電話", "得意職種", "許可番号", "許可年月日"]


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="366092")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def write_rows(ws, rows, with_hp=True):
    for c, h in enumerate(HEADERS, 1):
        ws.cell(row=1, column=c, value=h)
    for i, r in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=r["事業主名称"])
        ws.cell(row=i, column=3, value=r.get("派遣労働者数")).number_format = "#,##0"
        ws.cell(row=i, column=4, value=r.get("公開拠点数"))
        ws.cell(row=i, column=5, value=r.get("拠点数"))
        ws.cell(row=i, column=6, value=r.get("派遣先件数")).number_format = "#,##0"
        ws.cell(row=i, column=7, value=r.get("マージン率_平均")).number_format = '0.0"%"'
        ws.cell(row=i, column=8, value=r.get("派遣料金_平均")).number_format = "#,##0"
        ws.cell(row=i, column=9, value=r.get("賃金_平均")).number_format = "#,##0"
        hp = r.get("公式HP", "") if with_hp else ""
        if hp:
            c10 = ws.cell(row=i, column=10, value=hp)
            c10.hyperlink = hp
            c10.font = Font(color="0563C1", underline="single", name="Arial")
        else:
            ws.cell(row=i, column=10, value="")
        ws.cell(row=i, column=11, value=r.get("本社所在地", ""))
        ws.cell(row=i, column=12, value=r.get("電話", ""))
        ws.cell(row=i, column=13, value=(r.get("得意職種") or "")[:80])
        ws.cell(row=i, column=14, value=r.get("許可番号", ""))
        ws.cell(row=i, column=15, value=r.get("許可年月日", ""))
    widths = [6, 34, 12, 10, 8, 11, 12, 13, 12, 42, 42, 15, 26, 13, 13]
    for col, w in zip("ABCDEFGHIJKLMNO", widths):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:O{len(rows)+1}"
    style_header(ws, len(HEADERS))


def main():
    rows = json.load(open(JSON_PATH, encoding="utf-8"))
    withnum = [r for r in rows if r.get("派遣労働者数") is not None]

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "派遣数TOP500(HP付き)"
    write_rows(ws1, withnum[:500], with_hp=True)

    ws2 = wb.create_sheet("実数あり全社")
    write_rows(ws2, withnum, with_hp=True)

    wb.save(OUT_PATH)
    top = withnum[:500]
    got = len([r for r in top if r.get("公式HP")])
    print(f"出力: {OUT_PATH}")
    print(f"  シート1: TOP500 (HP {got}/500)")
    print(f"  シート2: 実数あり {len(withnum):,}社")


if __name__ == "__main__":
    main()
