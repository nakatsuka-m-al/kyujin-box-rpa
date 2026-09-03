# -*- coding: utf-8 -*-
"""新卒サービス企業の『中途向けサービス有無』をHPから判定し、Excel出力する。

入力: shinsotsu_candidates.json（会社→新卒サービス）, shinsotsu_services.json（サービス→運営会社/URL候補）
手順:
  1. 運営会社が未解決のサービス見出しのうち、サービスURL候補があるものはサイトのフッター/会社概要から運営会社を解決
  2. 各会社の公式HPを特定（サービスサイトの「運営会社」リンク → だめなら Yahoo/Bing 検索+検証）
  3. HP（トップ＋サービス/事業ページ）＋サービスサイトを読み、中途向けサービスの根拠語を抽出
     - 自社の採用ページ（募集要項/エントリー等）の行は除外して数える
  4. 判定: あり / 第二新卒のみ / なし / 不明(HP取得不可)
1社ごとに .shinsotsu_judge.json に保存（中断再開可）
"""
import json, os, re, sys, time, random
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from urllib.parse import urlparse
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright
from jinzai_common import UA_POOL, normalize_name, new_page, find_official_site, verify_site, is_bad_url
from shinsotsu_extract import clean_corp

BASE = "/Users/masakatsu/Desktop/kyujin_box_rpa"
CAND = os.path.join(BASE, "shinsotsu_candidates.json")
SVC = os.path.join(BASE, "shinsotsu_services.json")
PROG = os.path.join(BASE, ".shinsotsu_judge.json")
XLSX = os.path.join(BASE, "新卒特化サービス企業リスト_アライアンス候補.xlsx")
TXTDIR = os.path.join(BASE, ".shinsotsu_txt")
os.makedirs(TXTDIR, exist_ok=True)

CORP_RE = re.compile(r"((?:株式会社|有限会社|合同会社|一般社団法人)[\s　]?[\w一-龥ぁ-んァ-ヴー・＆&．.\-'’]{1,24}|[\w一-龥ぁ-んァ-ヴー・＆&．.\-'’]{1,24}[\s　]?(?:株式会社|有限会社|合同会社))")
COPY_RE = re.compile(r"(?:©|\(c\)|copyright|Copyright|COPYRIGHT)[^\n]{0,40}?" + CORP_RE.pattern)

# 中途向けサービスの根拠語（強）
STRONG = ["転職エージェント", "転職サイト", "転職支援", "転職サービス", "転職相談", "転職成功", "転職活動", "転職情報", "転職求人",
          "中途採用支援", "中途採用向け", "中途採用サービス", "中途人材", "中途紹介", "中途向け", "中途採用のご担当",
          "キャリア採用支援", "経験者採用支援", "ハイクラス", "ミドル層", "ミドル人材", "エグゼクティブ", "即戦力人材",
          "社会人向け", "ビジネスパーソン", "求人サイト", "人材派遣", "派遣サービス", "派遣事業", "労働者派遣事業", "スカウト型転職",
          "中途採用を支援", "中途採用の支援", "転職希望者", "キャリアアドバイザーが転職", "中途採用", "キャリア採用", "経験者採用", "転職"]
LICENSE_LINE = re.compile(r"許可番号|許可証|派\d{2}-\d+|\d{2}-ユ-\d+|労働者派遣事業許可|有料職業紹介事業許可")
# 中途を示すが弱い（第二新卒系）
SECOND = ["第二新卒", "既卒", "20代", "フリーター", "ニート", "若手社会人", "早期離職"]
# 自社採用ページの行とみなすマーカー（これが含まれる行は数えない）
SELF_RECRUIT = re.compile(r"中途採用(はこちら|情報|ページ|サイト|募集|エントリー|について|の方|をご希望|を希望)|キャリア採用(はこちら|情報|ページ|サイト|募集|エントリー)|募集要項|募集職種|募集中|エントリー|応募|求める人物|一緒に働|仲間を|私たちと|当社で働|社員インタビュー|福利厚生|選考フロー|選考プロセス|入社後|キャリアパス|RECRUIT|Recruit|JOIN US|採用情報|採用サイト|中途採用$|中途採用\s*$|新卒採用$|キャリア採用$|働く環境|社風|職場")
NEW_GRAD = ["新卒", "就活", "学生", "インターン", "内定", "27卒", "28卒", "26卒", "就職活動", "大学生", "理系学生", "体育会"]

SVC_PATHS = ["/service", "/services", "/business", "/service/", "/services/", "/business/", "/company", "/about", "/company/business", "/solution", "/product", "/products", "/lp", "/corporate"]


def fetch(pg, url, wait=900):
    try:
        r = pg.goto(url, timeout=20000)
        pg.wait_for_timeout(wait)
        if r and r.status >= 400:
            return "", [], []
        text = pg.inner_text("body")
        links = pg.evaluate("() => Array.from(document.querySelectorAll('a[href]')).map(a=>[a.href, (a.innerText||'').trim().slice(0,40)])")
        navs = pg.evaluate("() => Array.from(document.querySelectorAll('nav a, header a, [class*=menu] a, [class*=nav] a, [class*=gnav] a, footer a')).map(a=>(a.innerText||'').trim()).filter(t=>t&&t.length<30)")
        return text, links, navs
    except Exception:
        return "", [], []


def corp_from_site(pg, url):
    """サービスサイトからフッター等で運営会社と、会社HPリンクを取る"""
    text, links, _ = fetch(pg, url)
    corp, hp = "", ""
    if not text:
        return corp, hp, text
    tail = text[-3000:]
    m = COPY_RE.search(tail) or COPY_RE.search(text[:1500])
    if m:
        corp = m.group(1).strip()
    if not corp:
        m = re.search(r"(?:運営会社|運営|運営元|会社名|商号)[\s　]*[:：|｜\t\n]?[\s　]*" + CORP_RE.pattern, text[-4000:])
        if m:
            corp = m.group(1).strip()
    origin = "{0.scheme}://{0.netloc}".format(urlparse(url))
    # 運営会社/会社概要/コーポレートサイト リンク
    for href, t in links:
        if re.search(r"運営会社|会社概要|企業情報|コーポレート|運営元|会社情報|Company|COMPANY|Corporate|About", t or "") and href.startswith("http"):
            if urlparse(href).netloc != urlparse(url).netloc:
                hp = href; break
    if not hp:
        for href, t in links:
            if re.search(r"運営会社|会社概要|企業情報|運営元|会社情報", t or "") and href.startswith(origin):
                # 同一ドメインの会社概要ページ → そこから外部コーポレートサイトを辿る
                t2, l2, _ = fetch(pg, href, 600)
                if not corp and t2:
                    m = re.search(r"(?:会社名|商号|社名|運営会社)[\s　]*[:：|｜\t\n]?[\s　]*" + CORP_RE.pattern, t2)
                    if m: corp = m.group(1).strip()
                    m = COPY_RE.search(t2[-2000:])
                    if not corp and m: corp = m.group(1).strip()
                for h2, tt in l2:
                    if re.search(r"コーポレートサイト|会社HP|企業サイト|公式サイト|運営会社", tt or "") and h2.startswith("http") and urlparse(h2).netloc != urlparse(url).netloc:
                        hp = h2; break
                break
    corp = re.sub(r"\s+", "", corp)
    corp = re.sub(r"(?:All|ALL|all|Inc|Co|Ltd|Rights).*$", "", corp)
    corp = clean_corp(corp) if corp else ""
    return corp, hp, text


def analyze(texts, navs):
    """中途サービスの根拠を集計"""
    strong_hits, second_hits, newgrad_hits = {}, {}, 0
    snippets = []
    for text in texts:
        for line in text.split("\n"):
            l = line.strip()
            if not l or len(l) > 400:
                continue
            if SELF_RECRUIT.search(l) or LICENSE_LINE.search(l):
                continue
            for k in STRONG:
                if k in l:
                    strong_hits[k] = strong_hits.get(k, 0) + 1
                    if len(snippets) < 6 and all(k not in s for s in snippets):
                        snippets.append(l[:80])
            for k in SECOND:
                if k in l:
                    second_hits[k] = second_hits.get(k, 0) + 1
            for k in NEW_GRAD:
                if k in l:
                    newgrad_hits += 1
    nav_mid = [n for n in navs if re.search(r"転職|中途|キャリア採用支援|ハイクラス|ミドル|社会人", n) and not re.search(r"採用情報|中途採用$|キャリア採用$", n)]
    return strong_hits, second_hits, newgrad_hits, nav_mid, snippets


def judge(strong, second, newgrad, nav_mid, fetched):
    s_total0 = sum(strong.values())
    if not fetched and not (nav_mid or s_total0 >= 6):
        return "不明", "公式HP特定できず（サービスサイトのみでは中途なしと断定不可）"
    s_total = sum(strong.values())
    s_kinds = len(strong)
    reasons = []
    if nav_mid:
        reasons.append("メニューに中途系(" + "/".join(nav_mid[:3]) + ")")
    if s_kinds:
        reasons.append("中途語: " + ", ".join(f"{k}{v}" for k, v in sorted(strong.items(), key=lambda x: -x[1])[:6]))
    if second:
        reasons.append("第二新卒系語: " + ", ".join(f"{k}{v}" for k, v in second.items()))
    if (s_kinds >= 2 and s_total >= 3) or nav_mid or s_total >= 6:
        return "あり", "; ".join(reasons)
    if second and s_total <= 2:
        return "第二新卒のみ", "; ".join(reasons) or "第二新卒/既卒の記述のみ"
    if s_total >= 1:
        return "なし(要確認)", "; ".join(reasons) + "（中途語が少数のみ）"
    return "なし", "中途向けサービスの記述なし" + (f"（新卒語{newgrad}）" if newgrad else "")


def main(limit=None):
    cands = json.load(open(CAND, encoding="utf-8"))
    svcs = json.load(open(SVC, encoding="utf-8"))
    prog = json.load(open(PROG, encoding="utf-8")) if os.path.exists(PROG) else {}

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(user_agent=UA_POOL[0], locale="ja-JP")
        pg = ctx.new_page(); pg.set_default_timeout(20000)
        vp = new_page(browser); vp.set_default_timeout(15000)

        # ---- 1. 運営会社未解決サービスを、サイトから解決 ----
        svc_resolved = prog.setdefault("_svc_resolved", {})
        SKIP_HOST = re.compile(r"go\.jp|mhlw|affiliate|presco|link\.php|goo\.gl|a8\.net|moshimo|valuecommerce|accesstrade|click|\.pdf$")
        SKIP_HEAD = re.compile(r"[①-⑳]|質問|回答|とは|比較|おすすめ|について|なら|か$|い$|る$|た$|の$|ポイント|メリット|デメリット|注意|選び方|まとめ|一覧|表$|サイト$|サービス$|エージェント$|イベント$|コンテンツ|情報|特集|推移|機関")
        pend = [(k, v) for k, v in svcs.items() if not v["運営会社"] and v["URL候補"] and k not in svc_resolved
                and not SKIP_HEAD.search(v["サービス名"]) and not SKIP_HOST.search(v["URL候補"][0]) and len(v["サービス名"]) >= 2]
        print(f"運営会社未解決（URLあり）: {len(pend)}件", flush=True)
        for k, v in pend:
            corp, hp = "", ""
            for u in v["URL候補"][:2]:
                if is_bad_url(u):
                    continue
                corp, hp, _ = corp_from_site(pg, u)
                if corp:
                    break
            svc_resolved[k] = {"corp": corp, "hp": hp, "url": v["URL候補"][0]}
            if corp:
                ck = normalize_name(corp)
                c = cands.setdefault(ck, {"会社名": corp, "出現記事数": 0, "出典": [], "サービス名": [], "区分": []})
                if v["サービス名"] not in c["サービス名"]:
                    c["サービス名"].append(v["サービス名"])
                for s in v["出典"]:
                    if s not in c["出典"]:
                        c["出典"].append(s); c["出現記事数"] += 1
                for cc in v["区分"]:
                    if cc not in c["区分"]:
                        c["区分"].append(cc)
                c.setdefault("サービスURL", [])
                if v["URL候補"][0] not in c["サービスURL"]:
                    c["サービスURL"].append(v["URL候補"][0])
                if hp:
                    c["HPヒント"] = hp
                print(f"  解決: {v['サービス名'][:20]:<22} → {corp}", flush=True)
            json.dump(prog, open(PROG, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            json.dump(cands, open(CAND, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        # サービスURL/HPヒントを 既存会社にも付与
        for k, v in svcs.items():
            if v["運営会社"]:
                ck = normalize_name(v["運営会社"])
                if ck in cands and v["URL候補"]:
                    cands[ck].setdefault("サービスURL", [])
                    for u in v["URL候補"][:2]:
                        if u not in cands[ck]["サービスURL"] and len(cands[ck]["サービスURL"]) < 4:
                            cands[ck]["サービスURL"].append(u)
        json.dump(cands, open(CAND, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

        # ---- 2-3. 会社ごとにHP特定→判定 ----
        keys = sorted(cands.keys(), key=lambda k: -cands[k]["出現記事数"])
        if limit:
            keys = keys[:limit]
        t0 = time.time(); done = 0
        for k in keys:
            if k in prog:
                continue
            c = cands[k]
            name = c["会社名"]
            hp, how = "", ""
            svc_urls = [u for u in c.get("サービスURL", []) if not is_bad_url(u)]
            # a. サービスサイトのフッターから会社HP
            cand_hps = []
            if c.get("HPヒント"):
                cand_hps.append(c["HPヒント"])
            for u in svc_urls[:2]:
                corp2, hp2, _ = corp_from_site(pg, u)
                if hp2:
                    cand_hps.append(hp2)
            for h in cand_hps:
                try:
                    ok = verify_site(vp, h, name, "", "", browser)
                except Exception:
                    ok = ""
                if ok:
                    hp, how = ok, "サービスサイトの運営会社リンク"; break
            # c. 検索
            if not hp:
                try:
                    hit = find_official_site(browser, vp, name, "", "")
                except Exception:
                    hit = ""
                if hit and not is_bad_url(hit):
                    hp, how = hit, "検索で取得"
            # b. サービスサイト自体が会社サイトを兼ねているケース（会社名がタイトルに出る）
            if not hp:
                for u in svc_urls[:2]:
                    origin = "{0.scheme}://{0.netloc}/".format(urlparse(u))
                    try:
                        ok = verify_site(vp, origin, name, "", "", browser)
                    except Exception:
                        ok = ""
                    if ok:
                        hp, how = ok, "サービスサイト=会社サイト"; break
            # 判定用テキスト収集
            texts, navs = [], []
            fetched = False
            if hp:
                origin = "{0.scheme}://{0.netloc}".format(urlparse(hp))
                t, links, nv = fetch(pg, hp)
                if t:
                    fetched = True; texts.append(t); navs += nv
                    seen = set()
                    for href, tt in links:
                        if href.startswith(origin) and re.search(r"service|business|solution|product|jigyo|jigyou|事業|サービス|ソリューション", href + (tt or ""), re.I) and href not in seen and len(seen) < 6:
                            seen.add(href)
                    for href in list(seen) + [origin + sp for sp in SVC_PATHS]:
                        if len(texts) >= 8:
                            break
                        if href.rstrip("/") == hp.rstrip("/"):
                            continue
                        t2, _, nv2 = fetch(pg, href, 500)
                        if len(t2) > 200:
                            texts.append(t2); navs += nv2
            for u in svc_urls[:2]:
                t, _, nv = fetch(pg, u, 600)
                if t:
                    fetched = True; texts.append(t); navs += nv
            json.dump({"texts": texts, "navs": navs, "hp": hp}, open(os.path.join(TXTDIR, k[:60] + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
            strong, second, newgrad, nav_mid, snippets = analyze(texts, navs)
            cls, why = judge(strong, second, newgrad, nav_mid, fetched and bool(hp))
            prog[k] = {"name": name, "hp": hp, "how": how, "cls": cls, "why": why, "snip": snippets,
                       "svc_urls": svc_urls[:3], "pages": len(texts), "strong": strong, "second": second,
                       "newgrad": newgrad, "nav_mid": nav_mid, "fetched": fetched}
            json.dump(prog, open(PROG, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            done += 1
            el = time.time() - t0
            eta = (el / done) * (len([x for x in keys if x not in prog])) / 60
            print(f"[{done}] {name[:22]:<24} {cls:<9} {hp[:40] or '-':<42} 残り約{eta:.0f}分", flush=True)
        browser.close()
    write_excel(cands, prog)


def write_excel(cands, prog):
    wb = openpyxl.Workbook()
    sheets = {"中途なし(アライアンス候補)": [], "第二新卒のみ": [], "中途あり(除外)": [], "HP不明": []}
    for k, c in cands.items():
        v = prog.get(k)
        if not v:
            continue
        cls = v["cls"]
        row = [c["会社名"], v["hp"], "/".join(c.get("区分", [])), "、".join(c["サービス名"][:5]),
               "\n".join(v.get("svc_urls", [])[:2]), cls, v["why"][:300], "\n".join(v.get("snip", [])[:4]),
               c["出現記事数"], "\n".join(c["出典"][:2])]
        if cls.startswith("なし"):
            sheets["中途なし(アライアンス候補)"].append(row)
        elif cls == "第二新卒のみ":
            sheets["第二新卒のみ"].append(row)
        elif cls == "あり":
            sheets["中途あり(除外)"].append(row)
        else:
            sheets["HP不明"].append(row)
    header = ["会社名", "会社HP", "区分(A紹介/B媒体/C支援)", "新卒向けサービス名", "サービスURL", "中途サービス有無", "判定根拠", "HP抜粋", "掲載記事数", "出典記事"]
    widths = [28, 36, 12, 40, 40, 14, 50, 60, 9, 50]
    first = True
    for sname, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sname
        ws.append(header)
        for i, h in enumerate(header, 1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", start_color="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        rows.sort(key=lambda r: (-r[8], r[0]))
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            if row[1].value:
                row[1].hyperlink = row[1].value; row[1].font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(header))}{max(ws.max_row,2)}"
    # 説明シート
    ws = wb.create_sheet("説明")
    for line in [
        "■ 作り方", "新卒向けサービスの比較・まとめ記事（約70本）から、紹介されているサービス名と運営会社を抽出し、会社単位に統合。",
        "各社の公式HP（＋サービスサイト）をクロールし、中途向けサービス（転職エージェント/転職サイト/中途採用支援/ハイクラス/派遣 等）の記述の有無で判定。",
        "自社の採用ページ（募集要項・エントリー等）の行は除外して数えています。",
        "", "■ 判定", "なし＝中途向けサービスの記述なし（アライアンス候補）", "なし(要確認)＝中途語が少数のみ（ほぼ候補だが人の目で確認推奨）",
        "第二新卒のみ＝第二新卒/既卒向けの記述はあるが一般中途の記述なし", "あり＝中途向けサービスあり（除外）", "HP不明＝公式HPが特定できず未判定",
        "", "■ 注意", "判定は機械的なキーワード根拠です。「判定根拠」「HP抜粋」列を確認してください。",
        "区分 A=新卒紹介/就活エージェント、B=求人サイト/DR/逆求人/イベント/インターン、C=採用支援/RPO/コンサル/ツール（出典記事のカテゴリ）",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    wb.save(XLSX)
    from collections import Counter
    print("\n===== 完了 =====", dict(Counter(v["cls"] for k, v in prog.items() if not k.startswith("_"))))
    print("出力:", XLSX)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "excel":
        write_excel(json.load(open(CAND, encoding="utf-8")), json.load(open(PROG, encoding="utf-8")))
    else:
        main(int(sys.argv[1]) if len(sys.argv) > 1 else None)
