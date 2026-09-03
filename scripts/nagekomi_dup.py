# -*- coding: utf-8 -*-
"""② 投げ込み完了派遣の送信成功企業 → 人材紹介シートK列に「重複」"""
import sys, os, re, unicodedata, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from jinzai_common import normalize_name

SRC = "/Users/masakatsu/Desktop/kyujin_box_rpa/投げ込み_更新.xlsx"

def key(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip()
    s = re.sub(r"[\s　]+", "", s)
    s = re.sub(r"株式会社|有限会社|合同会社|（株）|\(株\)|㈱|（有）|\(有\)", "", s)
    s = s.upper()
    return s

def main(write=True):
    wb = openpyxl.load_workbook(SRC)
    ws = wb["人材紹介"]; wh = wb["投げ込み完了派遣"]
    sent = {}
    for r in range(2, wh.max_row + 1):
        if wh.cell(r, 3).value == "送信成功" and wh.cell(r, 1).value:
            sent[key(wh.cell(r, 1).value)] = wh.cell(r, 1).value
    # ドメインでも照合（URL列）
    sent_dom = {}
    for r in range(2, wh.max_row + 1):
        u = str(wh.cell(r, 2).value or "")
        m = re.search(r"https?://(?:www\.)?([^/]+)", u)
        if m and wh.cell(r, 3).value == "送信成功":
            sent_dom[m.group(1).lower()] = wh.cell(r, 1).value
    hits = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        k = key(name)
        hp = str(ws.cell(r, 5).value or "")
        m = re.search(r"https?://(?:www\.)?([^/]+)", hp)
        dom = m.group(1).lower() if m else None
        why = None
        if k in sent:
            why = f"社名一致:{sent[k]}"
        elif dom and dom in sent_dom:
            why = f"HPドメイン一致:{sent_dom[dom]}"
        if why:
            hits.append((r, name, why))
            if write:
                ws.cell(r, 11).value = "重複"
    for h in hits: print(h)
    print("重複", len(hits), "件 / 送信成功", len(sent))
    if write:
        wb.save(SRC)
    return hits

if __name__ == "__main__":
    main(write=(len(sys.argv) > 1 and sys.argv[1] == "write"))
