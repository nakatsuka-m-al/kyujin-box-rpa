# -*- coding: utf-8 -*-
"""フェーズ3: 企業リストJSONをExcelに出力する"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

JSON_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/top500_companies.json"
OUT_PATH = "/Users/masakatsu/Desktop/kyujin_box_rpa/人材紹介_正社員紹介数ランキング_全国TOP500.xlsx"

HEADERS = ["順位", "事業主名称", "正社員(無期)紹介数", "就職者数(4ヶ月以上・有期含む)",
           "公式HP", "本社所在地", "電話", "許可番号", "許可年月日"]


def main():
    with open(JSON_PATH, encoding="utf-8") as f:
        rows = json.load(f)
    rows = sorted(rows, key=lambda x: -x["正社員(無期)紹介数"])

    wb = Workbook()
    ws = wb.active
    ws.title = "正社員紹介数ランキング"

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="366092")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=r["事業主名称"])
        c3 = ws.cell(row=i, column=3, value=r["正社員(無期)紹介数"])
        c3.number_format = "#,##0"
        c4 = ws.cell(row=i, column=4, value=r.get("就職者数_4ヶ月以上"))
        c4.number_format = "#,##0"
        hp = r.get("公式HP", "")
        if hp:
            c5 = ws.cell(row=i, column=5, value=hp)
            c5.hyperlink = hp
            c5.font = Font(color="0563C1", underline="single", name="Arial")
        else:
            ws.cell(row=i, column=5, value="")
        ws.cell(row=i, column=6, value=r.get("本社所在地", ""))
        ws.cell(row=i, column=7, value=r.get("電話", ""))
        ws.cell(row=i, column=8, value=r.get("許可番号", ""))
        ws.cell(row=i, column=9, value=r.get("許可年月日", ""))

    for col, w in zip("ABCDEFGHI", [6, 38, 16, 18, 46, 46, 15, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(rows)+1}"
    ws.row_dimensions[1].height = 32

    wb.save(OUT_PATH)
    got = len([r for r in rows if r.get("公式HP")])
    print(f"出力: {OUT_PATH}")
    print(f"  {len(rows)}社 / HP有 {got}社 ({100*got//max(len(rows),1)}%)")


if __name__ == "__main__":
    main()
