# -*- coding: utf-8 -*-
"""戦略.xlsx の2シート（派遣数Top500 / 派遣拠点数Top500）の空欄「公式HP」を埋める。

方針:
- 間違ったURLを載せるくらいなら空欄のまま（検証に通ったものだけ記入）
- まず手元の全データソース（KBリスト解決結果・各種json）から許可番号で流用
- 残りをフルチェーン検索（Yahoo/Bing + 社名・住所・電話の複数クエリ）
- シート内の他のデータ・書式は一切変更しない（openpyxlで該当セルのみ書き込み）
"""
import json
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright
from jinzai_common import (
    check_maintenance_window, new_page, find_official_site, is_bad_url,
)

BASE_DIR = "/Users/masakatsu/Desktop/kyujin_box_rpa"
XLSX_IN = "/Users/masakatsu/Downloads/戦略.xlsx"
XLSX_OUT = os.path.join(BASE_DIR, "戦略_HP補完済み.xlsx")
PROGRESS = os.path.join(BASE_DIR, ".strategy_fill.json")


def collect_known_hps():
    """既存の全データソースから 許可番号->HP を集める"""
    hp = {}
    for path in ("haken_full.json", "haken_companies.json"):
        fp = os.path.join(BASE_DIR, path)
        if os.path.exists(fp):
            for r in json.load(open(fp, encoding="utf-8")):
                lic, u = r.get("許可番号"), r.get("公式HP")
                if lic and u and lic not in hp:
                    hp[lic] = u
    # KBリストの解決結果（stage Dで取得したHPを含む）
    fp = os.path.join(BASE_DIR, ".kb_resolution.json")
    if os.path.exists(fp):
        for v in json.load(open(fp, encoding="utf-8")).values():
            lic, u = v.get("派遣番号"), v.get("会社HP")
            if lic and u and lic not in hp:
                hp[lic] = u
    return hp


def main():
    check_maintenance_window()

    wb = openpyxl.load_workbook(XLSX_IN)
    known = collect_known_hps()
    print(f"既知HP: {len(known):,}件", flush=True)

    prog = {}
    if os.path.exists(PROGRESS):
        prog = json.load(open(PROGRESS, encoding="utf-8"))

    # 空欄セルの一覧を作る（シート名, 行, 会社名, 住所, 電話, 許可番号）
    targets = []
    for ws in wb.worksheets:
        header = [str(c.value) if c.value else "" for c in ws[1]]
        if "公式HP" not in header:
            continue
        hp_col = header.index("公式HP") + 1
        name_col = header.index("事業主名称") + 1
        addr_col = header.index("本社所在地") + 1 if "本社所在地" in header else None
        tel_col = header.index("電話") + 1 if "電話" in header else None
        lic_col = header.index("許可番号") + 1 if "許可番号" in header else None
        for row in range(2, ws.max_row + 1):
            name = ws.cell(row=row, column=name_col).value
            if not name:
                continue
            if ws.cell(row=row, column=hp_col).value:
                continue
            targets.append({
                "sheet": ws.title, "row": row, "hp_col": hp_col,
                "name": str(name).strip(),
                "addr": str(ws.cell(row=row, column=addr_col).value or "") if addr_col else "",
                "tel": str(ws.cell(row=row, column=tel_col).value or "") if tel_col else "",
                "lic": str(ws.cell(row=row, column=lic_col).value or "") if lic_col else "",
            })
    print(f"空欄セル: {len(targets)}件", flush=True)

    # 1) 許可番号で既知HPを流用
    resolved = {}   # lic or name -> url
    for t in targets:
        key = t["lic"] or t["name"]
        if key in prog:
            resolved[key] = prog[key]
        elif t["lic"] and t["lic"] in known:
            resolved[key] = known[t["lic"]]
    reused = len([1 for t in targets if resolved.get(t["lic"] or t["name"])])
    print(f"既存データ流用: {reused}件", flush=True)

    # 2) 残りを検索（許可番号で重複排除）
    todo = {}
    for t in targets:
        key = t["lic"] or t["name"]
        if not resolved.get(key):
            todo[key] = t
    print(f"検索対象: {len(todo)}社", flush=True)

    if todo:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            vp = new_page(browser)
            vp.set_default_timeout(15000)
            done = 0
            t0 = time.time()
            for key, t in todo.items():
                try:
                    hp = find_official_site(browser, vp, t["name"], t["addr"], t["tel"])
                except Exception:
                    hp = ""
                if hp and is_bad_url(hp):
                    hp = ""
                resolved[key] = hp
                prog[key] = hp
                json.dump(prog, open(PROGRESS, "w", encoding="utf-8"), ensure_ascii=False)
                done += 1
                el = time.time() - t0
                eta = (el / done) * (len(todo) - done) / 60
                print(f"[{done}/{len(todo)}] {t['name'][:24]:<26} -> "
                      f"{hp[:44] if hp else '(未検証のため空欄)'}  残り約{eta:.0f}分", flush=True)
            browser.close()

    # 3) 書き込み（該当セルのみ）
    filled = 0
    for t in targets:
        key = t["lic"] or t["name"]
        url = resolved.get(key, "")
        if url:
            ws = wb[t["sheet"]]
            c = ws.cell(row=t["row"], column=t["hp_col"], value=url)
            c.hyperlink = url
            c.font = Font(color="0563C1", underline="single", name="Arial")
            filled += 1
    wb.save(XLSX_OUT)

    print(f"\n===== 完了 =====", flush=True)
    print(f"空欄{len(targets)}件中 {filled}件を記入（残り{len(targets)-filled}件は検証不合格のため空欄維持）",
          flush=True)
    print(f"出力: {XLSX_OUT}", flush=True)


if __name__ == "__main__":
    main()
