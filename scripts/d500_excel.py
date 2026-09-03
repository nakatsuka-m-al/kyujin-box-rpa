# -*- coding: utf-8 -*-
"""投げ込み_更新.xlsx に「人材紹介_有期含む順」シートを追加。
- D順TOP500全社(既存342社はE/F/K列を既存シートから流用、新規158社は .d500.json から)
- K列: 派遣送信成功との重複(社名正規化+HPドメイン照合)
- L列: 既存「人材紹介」シートに掲載済みか
"""
import json, os, re, sys, unicodedata
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BASE = "/Users/masakatsu/Desktop/kyujin_box_rpa"
XLSX = os.path.join(BASE, "投げ込み_更新.xlsx")

def key(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip()
    s = re.sub(r"[\s　]+", "", s)
    s = re.sub(r"株式会社|有限会社|合同会社|（株）|\(株\)|㈱|（有）|\(有\)", "", s)
    return s.upper()

def dom(u):
    m = re.search(r"https?://(?:www\.)?([^/]+)", str(u or ""))
    return m.group(1).lower() if m else None

def main():
    d500 = json.load(open(os.path.join(BASE, "top500_yuki_companies.json"), encoding="utf-8"))
    prog = json.load(open(os.path.join(BASE, ".d500.json"), encoding="utf-8"))
    wb = openpyxl.load_workbook(XLSX)
    ws0 = wb["人材紹介"]; wh = wb["投げ込み完了派遣"]
    # 既存シート: 許可番号 -> (HP, form, 重複)
    exist = {}
    for r in range(2, ws0.max_row + 1):
        lic = str(ws0.cell(r, 9).value or "").strip()
        if lic:
            exist[lic] = (ws0.cell(r, 5).value or "", ws0.cell(r, 6).value or "", ws0.cell(r, 11).value or "")
    # 派遣送信成功
    sent_name, sent_dom = {}, {}
    for r in range(2, wh.max_row + 1):
        if wh.cell(r, 3).value == "送信成功" and wh.cell(r, 1).value:
            nm = wh.cell(r, 1).value
            sent_name[key(nm)] = nm
            d_ = dom(wh.cell(r, 2).value)
            if d_: sent_dom[d_] = nm
    if "人材紹介_有期含む順" in wb.sheetnames:
        del wb["人材紹介_有期含む順"]
    ws = wb.create_sheet("人材紹介_有期含む順")
    header = ["順位", "事業主名称", "正社員(無期)紹介数", "就職者数(4ヶ月以上・有期含む)",
              "公式HP", "問い合わせフォーム", "本社所在地", "電話", "許可番号", "許可年月日", "重複", "既存リスト"]
    ws.append(header)
    for i in range(1, len(header) + 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = PatternFill("solid", start_color="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n_dup = n_exist = n_hp = n_form = 0
    for x in d500:
        lic = x["許可番号"]
        hp, form, dup = "", "", ""
        in_exist = ""
        if lic in exist:
            hp, form, dup = exist[lic]
            in_exist = "掲載あり"
            n_exist += 1
        elif lic in prog:
            hp, form = prog[lic]["hp"], prog[lic]["form"]
        # 派遣重複(既存流用が空でも再判定)
        if not dup:
            if key(x["事業主名称"]) in sent_name:
                dup = "重複"
            elif dom(hp) and dom(hp) in sent_dom:
                dup = "重複"
        if dup: n_dup += 1
        if hp: n_hp += 1
        if form: n_form += 1
        row = [x["順位"], x["事業主名称"], x["正社員(無期)紹介数"], x["就職者数_4ヶ月以上"],
               hp, form, x.get("本社所在地", ""), x.get("電話", ""), lic, x.get("許可年月日", ""), dup, in_exist]
        ws.append(row)
        r = ws.max_row
        for col in (5, 6):
            v = ws.cell(r, col).value
            if v:
                ws.cell(r, col).hyperlink = v
                ws.cell(r, col).font = Font(color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHIJKL", [6, 32, 10, 12, 36, 40, 40, 14, 14, 11, 8, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    wb.save(XLSX)
    print(f"500社出力 / HP {n_hp} / フォーム {n_form} / 派遣重複 {n_dup} / 既存リスト掲載 {n_exist}")

if __name__ == "__main__":
    main()
