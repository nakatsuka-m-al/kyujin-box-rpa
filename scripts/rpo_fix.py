# -*- coding: utf-8 -*-
"""RPO.xlsx の E列（会社HP）を全件検索し直して正確なURLに置き換える。

- 既存URLは信用しない（生成された誤URLが混在）。全社ゼロから検索する
- ただし既存URLも「候補の1つ」として検証にかける（正しければそのまま採用）
- 検証に通らなければ空欄（誤URLを載せない）
- 「アデコ株式会社（LHH事業部）」のような括弧付きは法人名部分で検索する
- 元ファイルの他列・書式は触らず、E列のみ更新。別名で保存
"""
import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright
from jinzai_common import (
    check_maintenance_window, new_page, find_official_site, verify_site,
    is_bad_url, extract_city,
)

XLSX_IN = "/Users/masakatsu/Downloads/RPO.xlsx"
XLSX_OUT = "/Users/masakatsu/Desktop/kyujin_box_rpa/RPO_HP修正済み.xlsx"
PROGRESS = "/Users/masakatsu/Desktop/kyujin_box_rpa/.rpo_fix.json"


def clean_company(name):
    """検索用に法人名だけ取り出す（括弧内の事業部名などを除く）"""
    s = str(name).strip()
    s = re.sub(r"[（(][^）)]*[）)]", "", s)      # （LHH事業部）等
    s = re.sub(r"[\s　]+", " ", s).strip()
    return s


def main(limit=None):
    check_maintenance_window()

    wb = openpyxl.load_workbook(XLSX_IN)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(header)}
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=ci["会社名"]).value]
    if limit:
        rows = rows[:limit]
    print(f"対象: {len(rows)}社", flush=True)

    prog = json.load(open(PROGRESS, encoding="utf-8")) if os.path.exists(PROGRESS) else {}

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        vp = new_page(browser)
        vp.set_default_timeout(15000)
        t0 = time.time()
        done = 0
        for r in rows:
            raw = str(ws.cell(row=r, column=ci["会社名"]).value)
            pkey = f"{r}|{raw}"
            if pkey in prog or raw in prog:   # 旧キー(会社名のみ)との互換
                continue
            name = clean_company(raw)
            addr = str(ws.cell(row=r, column=ci["本社所在地"]).value or "")
            old = str(ws.cell(row=r, column=ci["会社HP"]).value or "").strip()
            _, city = extract_city(addr)

            new_url, how = "", ""
            # 1) 既存URLが本物なら採用（検証はフルに通す）
            if old and old.startswith("http") and not is_bad_url(old):
                ok = ""
                for attempt in range(2):          # 一時的な403等に備え1回リトライ
                    try:
                        ok = verify_site(vp, old, name, city, "", browser)
                    except Exception:
                        ok = ""
                    if ok:
                        break
                    time.sleep(4)
                if ok:
                    new_url, how = ok, "既存URL検証OK"
            # 2) 検証に通らなければ検索し直し
            if not new_url:
                try:
                    hit = find_official_site(browser, vp, name, addr, "")
                except Exception:
                    hit = ""
                if hit and not is_bad_url(hit):
                    new_url, how = hit, "再検索で取得"
                else:
                    how = "見つからず(空欄)"

            prog[pkey] = {"new": new_url, "old": old, "how": how}
            json.dump(prog, open(PROGRESS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            done += 1
            el = time.time() - t0
            eta = (el / done) * (len(rows) - done) / 60
            chg = "" if new_url == old.rstrip("/") or new_url == old else " ←変更"
            print(f"[{done}/{len(rows)}] {name[:22]:<24} {how:<12} {new_url[:44] or '-'}{chg}  残り約{eta:.0f}分", flush=True)
        browser.close()

    # 書き込み
    changed = filled = blank = 0
    for r in rows:
        raw = str(ws.cell(row=r, column=ci["会社名"]).value)
        v = prog.get(f"{r}|{raw}") or prog.get(raw)
        if not v:
            continue
        cell = ws.cell(row=r, column=ci["会社HP"])
        if v["new"]:
            if v["new"].rstrip("/") != v["old"].rstrip("/"):
                changed += 1
            cell.value = v["new"]
            cell.hyperlink = v["new"]
            cell.font = Font(color="0563C1", underline="single")
            filled += 1
        else:
            cell.value = ""
            cell.hyperlink = None
            blank += 1
    wb.save(XLSX_OUT)
    print(f"\n===== 完了 =====")
    print(f"HP確定 {filled}社（うち修正{changed}社） / 特定できず空欄 {blank}社")
    print(f"出力: {XLSX_OUT}")


if __name__ == "__main__":
    lim = int(sys.argv[1]) if len(sys.argv) > 1 else None
    main(lim)
